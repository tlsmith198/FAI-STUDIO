"""Importers for parts lists, notes lists and measured-results files.

Supported tabular sources: .csv .tsv .txt .xlsx .xlsm .docx (tables) .pdf (table detection)
Supported notes sources:   .txt .docx .pdf .csv .xlsx (one note per row)
"""
from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import Optional

TABLE_EXT = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm", ".docx", ".pdf"}

COMPONENT_FIELDS = {
    "item_no": ["ITEM", "ITEM NO", "ITEM #", "FIND", "FIND NO", "FIND NUMBER", "LINE", "NO", "#", "BALLOON"],
    "part_number": ["PART NUMBER", "PART NO", "PART #", "P/N", "PN", "PART", "IDENTIFYING NUMBER",
                    "CAGE/PART", "COMPONENT", "DWG NO", "DRAWING NUMBER", "ITEM NUMBER"],
    "part_name": ["DESCRIPTION", "NOMENCLATURE", "PART NAME", "NAME", "NOMENCLATURE OR DESCRIPTION", "DESC"],
    "qty": ["QTY", "QUANTITY", "QTY REQD", "QTY REQ", "QTY/ASSY", "Q/A", "REQD"],
    "serial_number": ["SERIAL", "SERIAL NUMBER", "S/N", "SN"],
    "fai_report_number": ["FAI", "FAI REPORT", "FAI REPORT NUMBER", "FAI NO"],
    "notes": ["MATERIAL", "SPEC", "SPECIFICATION", "MATERIAL/SPEC", "NOTES", "REMARKS", "CAGE", "CAGE CODE"],
}
MATERIAL_FIELDS = {
    "name": ["MATERIAL", "DESCRIPTION", "PROCESS", "NAME", "NOMENCLATURE", "MATERIAL/PROCESS", "PART NAME"],
    "spec_number": ["SPEC", "SPECIFICATION", "SPEC NUMBER", "SPECIFICATION NUMBER", "MATERIAL SPEC", "PART NUMBER", "P/N"],
    "code": ["CODE", "CLASS", "TYPE", "GRADE", "CONDITION", "TEMPER", "FORM"],
    "supplier": ["SUPPLIER", "VENDOR", "SOURCE", "MFR", "MANUFACTURER"],
    "supplier_code": ["SUPPLIER CODE", "CAGE", "CAGE CODE", "VENDOR CODE"],
    "coc_number": ["COC", "C OF C", "CERT", "CERT NO", "CERTIFICATE", "CERT #", "HEAT", "HEAT NO", "LOT"],
    "comments": ["NOTES", "REMARKS", "COMMENTS", "QTY"],
}
RESULT_FIELDS = {
    "char_no": ["CHAR", "CHAR NO", "CHAR #", "CHARACTERISTIC", "BALLOON", "BALLOON NO", "FEATURE", "NO", "#", "ID", "ITEM"],
    "results": ["ACTUAL", "MEASURED", "RESULT", "RESULTS", "VALUE", "MEAS", "MEASUREMENT", "ACT"],
    "tooling": ["GAGE", "GAUGE", "TOOL", "TOOLING", "INSTRUMENT", "EQUIPMENT"],
}


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def read_table(path: str | Path) -> list[list[str]]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        rows = []
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True):
                cells = [_clean(c) for c in r]
                if any(cells):
                    rows.append(cells)
            if rows:
                break
        return rows
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        rows = []
        for t in d.tables:
            for r in t.rows:
                cells = [_clean(c.text) for c in r.cells]
                if any(cells):
                    rows.append(cells)
        return rows
    if ext == ".pdf":
        import pymupdf
        rows = []
        with pymupdf.open(path) as doc:
            for page in doc:
                try:
                    tabs = page.find_tables()
                except Exception:
                    continue
                for t in tabs.tables:
                    for r in t.extract():
                        cells = [_clean(c) for c in r]
                        if any(cells):
                            rows.append(cells)
        return rows
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel_tab if "\t" in sample else csv.excel
    rows = [[_clean(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return [r for r in rows if any(r)]


def _norm(h: str) -> str:
    return re.sub(r"[^A-Z0-9#/ ]", "", h.upper()).strip()


def detect_mapping(rows: list[list[str]], fields: dict[str, list[str]]) -> tuple[int, dict[str, int]]:
    """Find the header row and map target field -> column index."""
    best = (-1, {}, 0)
    for idx, row in enumerate(rows[:25]):
        mapping: dict[str, int] = {}
        score = 0
        used = set()
        for field, aliases in fields.items():
            for ci, cell in enumerate(row):
                if ci in used:
                    continue
                n = _norm(cell)
                if not n:
                    continue
                if n in aliases:
                    mapping[field] = ci
                    used.add(ci)
                    score += 2
                    break
            else:
                for ci, cell in enumerate(row):
                    n = _norm(cell)
                    if ci in used or not n or len(n) > 40:
                        continue
                    if any(n.startswith(a) or a in n.split() for a in aliases if len(a) > 2):
                        mapping[field] = ci
                        used.add(ci)
                        score += 1
                        break
        if score > best[2]:
            best = (idx, mapping, score)
    return best[0], best[1]


def apply_mapping(rows: list[list[str]], header_idx: int, mapping: dict[str, int]) -> list[dict]:
    out = []
    for r in rows[header_idx + 1:]:
        rec = {f: (r[c] if c is not None and 0 <= c < len(r) else "") for f, c in mapping.items()}
        if any(v for v in rec.values()):
            out.append(rec)
    return out


# ---------------------------------------------------------------------------- notes
NOTE_RE = re.compile(r"^\s*(?:NOTE\s*)?(\d{1,3})\s*[.)]\s+(.+)$", re.I)
SUB_RE = re.compile(r"^\s*([A-Z])\s*[.)]\s+(.+)$")
# title block / banner text that means the notes block has ended (NOT "UNLESS OTHERWISE NOTED", which
# appears inside real notes)
END_NOTES_RE = re.compile(r"^(DWG\.?\s*NO|DRAWING\s+NO|SCALE\b|CAGE\b|TITLE:|DRAWN\b|CHECKED\b|APPROVED\b)|"
                          r"\bSHEET\s+\d+\s+OF\s+\d+|PROPRIETARY|EXPORT\s+CONTROL|DISTRIBUTION\s+STATEMENT|"
                          r"DESTRUCTION\s+NOTICE|WARNING\s*[-:]\s*THIS\s+DOCUMENT|STATE:\s*RELEASED|"
                          r"^REV(ISION)?\s*[:.]?\s*[A-Z0-9-]{1,3}\b.*\bSHEET\b", re.I)
HEADER_RE = re.compile(r"^\s*(GENERAL\s+)?NOTES?(\s*[-,:]?\s*UNLESS\s+OTHERWISE\s+(SPECIFIED|NOTED|STATED))?"
                       r"\s*:?\s*(\(.*\))?\s*$", re.I)
CONT_ON_RE = re.compile(r"\bNOTES?\s+CONT[\w'’]*\.?\s+(ON|TO)\s+(SHEET|SH\.?|NEXT)", re.I)
CONT_FROM_RE = re.compile(r"\bNOTES?\s+CONT[\w'’]*\.?\s+FROM\b", re.I)


def read_notes_text(path: str | Path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            text = "\n".join(p.get_text() for p in doc)
            if len(re.findall(r"\d+\.\s+[A-Z]", text)) < 2:
                text = _ocr_pdf_text(doc, path) or text          # scanned notes list
        return text
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" ".join(c.text for c in r.cells))
        # Word auto-numbered lists lose their numbers: number them if none present
        if not any(NOTE_RE.match(p) for p in parts if p.strip()):
            parts = [f"{i}. {p}" for i, p in enumerate([p for p in parts if p.strip() and not HEADER_RE.match(p)], 1)]
        return "\n".join(parts)
    if ext in (".xlsx", ".xlsm", ".csv", ".tsv"):
        rows = read_table(path)
        lines = []
        for i, r in enumerate(rows, 1):
            cells = [c for c in r if c]
            if not cells:
                continue
            if len(cells) >= 2 and re.fullmatch(r"\d{1,3}[.)]?", cells[0]):
                lines.append(f"{cells[0].rstrip('.)')}. {' '.join(cells[1:])}")
            elif NOTE_RE.match(cells[0]):
                lines.append(" ".join(cells))
            elif not HEADER_RE.match(" ".join(cells)):
                lines.append(f"{i}. {' '.join(cells)}")
        return "\n".join(lines)
    return path.read_text(encoding="utf-8-sig", errors="replace")


def split_inline_subparagraphs(text: str) -> str:
    """'A. PRESS-FIT … B. PRESS-FIT …' on one line -> one line per lettered sub-paragraph (only when the
    letters follow on in sequence, so 'SEE FIG. B. ' style text is left alone)."""
    lines = text.split("\n")
    out = []
    for ln in lines:
        m = re.match(r"^([A-Z])\.\s", ln)
        if not m:
            out.append(ln)
            continue
        nxt = chr(ord(m.group(1)) + 1)
        while True:
            k = re.search(rf"(?<=[.)])\s+{nxt}\.\s+(?=[A-Z0-9])", ln)
            if not k:
                break
            out.append(ln[:k.start()].rstrip())
            ln = ln[k.start():].lstrip()
            nxt = chr(ord(nxt) + 1)
        out.append(ln)
    return "\n".join(out)


def parse_notes(text: str) -> list[dict]:
    """Return [{'label': '3', 'text': '...'}, …].

    Lettered sub-paragraphs ("A. …", "B. …") stay inside their note (on their own lines). A label written as
    "7A" is its own note (SpaceX style, flagged individually). Notes continue across sheets ("NOTES CONTINUED
    ON SHEET 2" … "NOTES CONTINUED FROM SHEET 1:"), and a note number may sit alone on the line before its text."""
    lines = text.splitlines()
    # If a NOTES header exists, start after it (drawing PDFs contain lots of other text)
    for i, ln in enumerate(lines):
        if HEADER_RE.match(ln):
            lines = lines[i + 1:]
            break
    notes: list[dict] = []
    current: Optional[dict] = None
    expected = 1
    paused = False
    pending: Optional[int] = None

    def ok(num):
        return not notes or num in (expected, expected + 1, expected + 2)

    def start(num, body, label=None):
        nonlocal current, expected
        current = {"label": label or str(num), "text": body.strip()}
        notes.append(current)
        if label is None:
            expected = num + 1

    flat = [ln.strip() for ln in lines if ln.strip()]
    for idx, s in enumerate(flat):
        mc = CONT_ON_RE.search(s)
        if mc:
            before = s[:mc.start()].strip()
            if before and current is not None and not paused:
                current["text"] += " " + before
            paused = True                               # resume at "NOTES CONTINUED FROM SHEET n"
            pending = None
            continue
        if CONT_FROM_RE.search(s) or HEADER_RE.match(s):
            paused = False
            continue
        if paused:
            continue
        if pending is not None:                         # "21" on its own line, text on this one
            start(pending, s)
            pending = None
            continue
        m = NOTE_RE.match(s)
        if m and ok(int(m.group(1))):
            start(int(m.group(1)), m.group(2))
            continue
        mn = re.fullmatch(r"(?:NOTE\s*)?(\d{1,3})\s*[.)]?", s, re.I)
        if mn and ok(int(mn.group(1))) and (current is None or len(current["text"]) > 3):
            pending = int(mn.group(1))
            continue
        m7 = re.match(r"^(\d{1,3})([A-Z])\s*[.)]?\s+(.+)$", s)
        if m7 and current is not None and re.match(r"\d+", current["label"]).group(0) == m7.group(1):
            start(int(m7.group(1)), m7.group(3), label=f"{m7.group(1)}{m7.group(2)}")
            continue
        sm = SUB_RE.match(s)
        if sm and current is not None:
            current["text"] += f"\n{sm.group(1)}. {sm.group(2).strip()}"
            continue
        if current is not None:
            if END_NOTES_RE.search(s) or len(current["text"]) > 8000:
                paused = True                           # title block / banner: the notes on this sheet are over
                continue
            # a finished sentence followed (within two lines) by title-block text: this line is title block too
            if current["text"].rstrip().endswith(".") and any(END_NOTES_RE.search(x) for x in flat[idx + 1:idx + 3]) \
                    and not NOTE_RE.match(s):
                paused = True
                continue
            current["text"] += " " + s
    for n in notes:
        n["text"] = split_inline_subparagraphs(
            "\n".join(re.sub(r"\s+", " ", part).strip() for part in n["text"].split("\n")).strip())
    return notes


def _ocr_pdf_text(doc, path) -> str:
    """OCR a scanned PDF into reading-order text lines (used for scanned notes lists)."""
    try:
        from . import ocr
        if not ocr.available():
            return ""
        import hashlib
        h = hashlib.sha256(Path(path).read_bytes()).hexdigest()
        out = []
        for page in doc:
            boxes = ocr.ocr_pdf_page(h, page, 200)
            boxes.sort(key=lambda b: (round(b[1] / 8), b[0]))
            rows: list[list] = []
            for b in boxes:
                if rows and abs(b[1] - rows[-1][0][1]) < 0.6 * (b[3] - b[1]):
                    rows[-1].append(b)
                else:
                    rows.append([b])
            out += [" ".join(x[4] for x in sorted(r, key=lambda x: x[0])) for r in rows]
        return "\n".join(out)
    except Exception:
        return ""
