"""Excel export of the report (AS9102 Forms 1-3 or Inspection Report)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CHAR = "080B0B"
ORANGE = "FF6A00"
thin = Side(style="thin", color="808080")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=CHAR)
REJ_FILL = PatternFill("solid", fgColor="FDE2E1")
ACC_FONT = Font(color="146C2E", bold=True)
REJ_FONT = Font(color="B42318", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def _title(ws, text, meta, rep, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, f"{meta.get('company_name', '')}  —  CAGE {meta.get('cage_code', '')}")
    c.font = Font(bold=True, size=12, color="FFFFFF")
    c.fill = HDR_FILL
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(2, 1, text)
    c.font = Font(bold=True, size=12, color=ORANGE)
    ws.row_dimensions[1].height = 20


def _fields(ws, row, pairs, ncols):
    """pairs: list of (label, value) laid out 4 per row (label above value)."""
    per = 4
    span = max(1, ncols // per)
    for i in range(0, len(pairs), per):
        for j, (label, value) in enumerate(pairs[i:i + per]):
            col = 1 + j * span
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
            ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + span - 1)
            lc = ws.cell(row, col, label)
            lc.font = Font(size=8, color="4A4F52")
            vc = ws.cell(row + 1, col, value)
            vc.font = Font(bold=True)
            vc.alignment = WRAP
            for rr in (row, row + 1):
                for cc in range(col, col + span):
                    ws.cell(rr, cc).border = BORDER
        row += 2
    return row + 1


def _table(ws, row, headers, rows, widths, disp_col=None):
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row, ci, h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDER
    ws.row_dimensions[row].height = 30
    for ri, r in enumerate(rows, row + 1):
        for ci, v in enumerate(r, 1):
            c = ws.cell(ri, ci, v)
            c.alignment = WRAP
            c.border = BORDER
        if disp_col is not None:
            d = r[disp_col]
            if d == "Reject":
                for ci in range(1, len(headers) + 1):
                    ws.cell(ri, ci).fill = REJ_FILL
                ws.cell(ri, disp_col + 1).font = REJ_FONT
            elif d == "Accept":
                ws.cell(ri, disp_col + 1).font = ACC_FONT
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = ws.cell(row + 1, 1)
    return row + len(rows) + 2


def export_excel(db, report_id: int, out_path: str | Path, report_type: str | None = None) -> Path:
    rep = db.get_report(report_id)
    if report_type:
        rep = dict(rep, report_type=report_type)          # export as another document type
    meta = db.get_meta()
    chars = db.rows("characteristics", report_id)
    comps = db.rows("components", report_id)
    mats = db.rows("materials", report_id)
    wb = Workbook()
    ident = [("1. Part Number", rep["part_number"]), ("2. Part Name", rep["part_name"]),
             ("3. Serial Number", rep["serial_number"]), ("4. FAI Report Number", rep["report_number"])]

    from .report_types import STANDARD, is_fai
    if is_fai(rep["report_type"]):
        ws = wb.active
        ws.title = "Form 1"
        _title(ws, "AS9102 Form 1: Part Number Accountability", meta, rep, 8)
        row = _fields(ws, 4, ident + [
            ("5. Part Revision Level", rep["part_rev"]), ("6. Drawing Number", rep["drawing_number"]),
            ("7. Drawing Revision Level", rep["drawing_rev"]), ("8. Additional Changes", rep["additional_changes"]),
            ("9. Manufacturing Process Reference", rep["mfg_process_ref"]),
            ("10. Organization Name", rep["organization_name"] or meta.get("company_name")),
            ("11. Supplier Code", rep["supplier_code"] or meta.get("cage_code")),
            ("12. Purchase Order Number", rep["po_number"]),
            ("13. Detail / Assembly FAI", rep["detail_or_assembly"]), ("14. Full / Partial FAI", rep["fai_type"]),
            ("Baseline P/N incl. Rev (Partial)", rep["partial_baseline"]), ("Reason for Partial FAI", rep["partial_reason"]),
        ], 8)
        row = _table(ws, row, ["Item", "15. Part Number", "16. Part Name", "Qty", "17. Part Serial Number",
                               "18. FAI Report Number", "Notes", ""],
                     [[c["item_no"], c["part_number"], c["part_name"], c["qty"], c["serial_number"],
                       c["fai_report_number"], c["notes"], ""] for c in comps],
                     [8, 22, 30, 7, 18, 18, 30, 4])
        has_nc = any(c["disposition"] == "Reject" or c["nonconformance"] for c in chars)
        _fields(ws, row, [("19. Contains documented nonconformance(s)?", "Yes" if has_nc else "No"),
                          ("20. FAI Complete Signature", rep["prepared_by"]), ("21. Date", rep["prepared_date"]),
                          ("22. Reviewed By", rep["reviewed_by"]), ("23. Date", rep["reviewed_date"]),
                          ("24. Customer Approval", rep["customer_approval"]),
                          ("25. Date", rep["customer_approval_date"]), ("26. Comments", rep["comments"])], 8)

        ws = wb.create_sheet("Form 2")
        _title(ws, "AS9102 Form 2: Product Accountability — Materials, Special Processes, Functional Testing", meta, rep, 8)
        row = _fields(ws, 4, ident, 8)
        _table(ws, row, ["Type", "5. Material or Process Name", "6. Specification Number", "7. Code",
                         "8. Special Process Supplier Code", "9. Customer Approval Verification",
                         "10. Certificate of Conformance / 12. Acceptance Report No.", "Comments"],
               [[m["category"], m["name"], m["spec_number"], m["code"],
                 " / ".join(x for x in (m["supplier"], m["supplier_code"]) if x), m["customer_approval"],
                 m["coc_number"], m["comments"]] for m in mats],
               [16, 30, 24, 12, 22, 14, 24, 30])

        ws = wb.create_sheet("Form 3")
        _title(ws, "AS9102 Form 3: Characteristic Accountability, Verification and Compatibility Evaluation",
               meta, rep, 13)
        row = _fields(ws, 4, ident, 12)
    else:
        ws = wb.active
        ws.title = "Inspection Report"
        _title(ws, "Inspection Report" if rep["report_type"] == STANDARD else "Dimensional Inspection Report",
               meta, rep, 13)
        row = _fields(ws, 4, [("Report Number", rep["report_number"]), ("Part Number", rep["part_number"]),
                              ("Part Name", rep["part_name"]), ("Revision", rep["part_rev"]),
                              ("Drawing / Rev", f"{rep['drawing_number']} {rep['drawing_rev']}"),
                              ("Customer", rep["customer"]), ("PO", rep["po_number"]), ("Job", rep["job_number"]),
                              ("Serial / Lot", _serial_text(rep)), ("Qty / Sample", f"{rep['quantity']} / {rep['sample_size']}"),
                              ("Inspector", rep["prepared_by"]), ("Date", rep["prepared_date"])], 12)

    _table(ws, row, ["5. Char No.", "6. Ref Location", "7. Designator", "Type", "8. Requirement", "Nominal",
                     "+Tol", "-Tol", "LSL", "USL", "9. Results", "Accept/Reject", "10. Tooling",
                     "11. NCR Number", "12. Comments"],
           [[c["char_no"], c["ref_location"], c["designator"], c["char_type"], c["requirement"], c["nominal"],
             c["tol_plus"], c["tol_minus"], c["lower_limit"], c["upper_limit"], c["results"], c["disposition"],
             c["tooling"], c["nonconformance"], c["comments"]] for c in chars],
           [8, 11, 11, 12, 34, 10, 8, 8, 10, 10, 16, 11, 14, 12, 30], disp_col=11)
    from . import multipart as mp
    if mp.applies(rep):
        _per_part_sheet(wb, rep, chars, meta)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.gridLines = False
    out_path = Path(out_path)
    wb.save(out_path)
    return out_path


def _serial_text(rep) -> str:
    from . import multipart as mp
    if not mp.applies(rep) or mp.part_count(rep) <= 1:
        return rep["serial_number"]
    sns = [x for x in mp.serials(rep) if x]
    return f"{mp.part_count(rep)} parts" + (": " + ", ".join(sns) if sns else "")


def _per_part_sheet(wb, rep, chars, meta):
    """One column per part (Low / High pair per part in low & high mode) + per-part summary."""
    from . import multipart as mp
    n, mm = mp.part_count(rep), mp.minmax(rep)
    sns = mp.serials(rep)
    ws = wb.create_sheet("Per-Part Readings")
    heads = ["Char No.", "Requirement", "LSL", "USL"]
    for i in range(1, n + 1):
        lab = f"P{i}" + (f" {sns[i - 1]}" if sns[i - 1] else "")
        heads += [f"{lab} Low", f"{lab} High"] if mm else [lab]
    heads += ["All Parts Results", "All Parts Accept/Reject"]
    ncols = len(heads)
    _title(ws, f"Per-Part Readings — {n} parts, " + ("low & high per part" if mm else "one reading per part"),
           meta, rep, ncols)
    rows, rej_cells = [], []
    for ri, c in enumerate(chars):
        rd = mp.parse(c.get("part_readings"))
        row = [c["char_no"], c["requirement"], c["lower_limit"], c["upper_limit"]]
        for i in range(1, n + 1):
            x = rd.get(i) or {}
            for v in ((x.get("lo", ""), x.get("hi", "")) if mm else (x.get("lo", ""),)):
                if mp.value_rejected(v, c, x.get("d", "")):
                    rej_cells.append((ri, len(row)))
                row.append(v)
        row += [c["results"], c["disposition"]]
        rows.append(row)
    start = 4
    _table(ws, start, heads, rows, [8, 34, 10, 10] + [10] * (ncols - 6) + [22, 12], disp_col=ncols - 1)
    for ri, ci in rej_cells:
        cell = ws.cell(start + 1 + ri, ci + 1)
        cell.font = REJ_FONT
        cell.fill = REJ_FILL
